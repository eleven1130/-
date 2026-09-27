"""生成「校园工具箱」双创大赛 PPT"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# 配色
BG_DARK = RGBColor(0x1A, 0x1A, 0x2E)       # 深蓝紫背景
ACCENT = RGBColor(0x00, 0xD2, 0xFF)         # 青色强调
ACCENT2 = RGBColor(0x7C, 0x4D, 0xFF)        # 紫色
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xCC, 0xCC, 0xDD)
GRAY = RGBColor(0x99, 0x99, 0xAA)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

blank = prs.slide_layouts[6]  # 空白


def set_bg(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    bg.shadow.inherit = False
    return bg


def add_text(slide, left, top, width, height, text, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = '微软雅黑'
    return tb


def add_rect(slide, left, top, width, height, color, alpha=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_accent_bar(slide, left, top, width=Inches(0.12), height=Inches(0.5), color=ACCENT):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    return bar


def page_number(slide, n, total):
    add_text(slide, Inches(12.3), Inches(7.0), Inches(1), Inches(0.4),
             f"{n} / {total}", size=11, color=GRAY, align=PP_ALIGN.RIGHT)


TOTAL = 11

# ============ Slide 1: 封面 ============
s = prs.slides.add_slide(blank)
set_bg(s, BG_DARK)
# 装饰圆
c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.5), Inches(-1.5), Inches(5), Inches(5))
c.fill.solid(); c.fill.fore_color.rgb = ACCENT2; c.line.fill.background()
c.fill.fore_color.brightness = 0.3
c2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-2), Inches(5), Inches(4), Inches(4))
c2.fill.solid(); c2.fill.fore_color.rgb = ACCENT; c2.line.fill.background()
c2.fill.fore_color.brightness = 0.4

add_text(s, Inches(0.8), Inches(2.0), Inches(10), Inches(0.6),
         "INNOVATION & ENTREPRENEURSHIP", size=14, color=ACCENT, bold=True)
add_text(s, Inches(0.8), Inches(2.7), Inches(11), Inches(1.5),
         "校园工具箱", size=60, color=WHITE, bold=True)
add_text(s, Inches(0.8), Inches(4.0), Inches(11), Inches(0.8),
         "Campus Tools — 让校园生活更高效", size=24, color=LIGHT)
add_text(s, Inches(0.8), Inches(5.2), Inches(11), Inches(0.5),
         "基于 it-tools 开源项目改造 · 课表空闲时段计算", size=16, color=GRAY)
add_text(s, Inches(0.8), Inches(6.5), Inches(11), Inches(0.4),
         "参赛团队：校园工具箱小组", size=13, color=GRAY)
page_number(s, 1, TOTAL)

# ============ Slide 2: 项目背景 ============
s = prs.slides.add_slide(blank)
set_bg(s, BG_DARK)
add_accent_bar(s, Inches(0.6), Inches(0.7))
add_text(s, Inches(0.9), Inches(0.6), Inches(10), Inches(0.7),
         "01  项目背景", size=32, color=WHITE, bold=True)
add_text(s, Inches(0.9), Inches(1.5), Inches(11), Inches(0.5),
         "为什么做「校园工具箱」？", size=18, color=ACCENT)

# 三个数据卡片
cards = [
    ("90%+", "大学生每天使用\n5 个以上效率工具"),
    ("∞", "工具分散在各个 App、\n网页、小程序中"),
    ("0", "专为大学生设计的\n聚合工具平台"),
]
for i, (num, desc) in enumerate(cards):
    left = Inches(0.8 + i * 4.1)
    add_rect(s, left, Inches(2.4), Inches(3.7), Inches(2.6), RGBColor(0x25, 0x25, 0x40))
    add_text(s, left, Inches(2.7), Inches(3.7), Inches(1.2),
             num, size=48, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, left + Inches(0.3), Inches(4.0), Inches(3.1), Inches(0.9),
             desc, size=14, color=LIGHT, align=PP_ALIGN.CENTER)

add_text(s, Inches(0.9), Inches(5.5), Inches(11.5), Inches(1.2),
         "it-tools 是一个成熟的开源工具箱（Vue 3 + TS，月活高，GPL-3.0），架构上「一个工具 = 一个组件」，\n"
         "新增工具门槛极低。我们以此为底座，将其改造为面向大学生的「校园工具箱」，\n"
         "在保留通用工具的同时，新增校园专属工具，并实现全中文界面与离线可用。",
         size=15, color=LIGHT)
page_number(s, 2, TOTAL)

# ============ Slide 3: 痛点分析 ============
s = prs.slides.add_slide(blank)
set_bg(s, BG_DARK)
add_accent_bar(s, Inches(0.6), Inches(0.7))
add_text(s, Inches(0.9), Inches(0.6), Inches(10), Inches(0.7),
         "02  痛点分析", size=32, color=WHITE, bold=True)
add_text(s, Inches(0.9), Inches(1.5), Inches(11), Inches(0.5),
         "「想约同学，却不知道大家什么时候都有空」", size=18, color=ACCENT)

pain_points = [
    ("⏳", "效率低", "手动对照多人课表找共同空闲，\n反复沟通「你周一有空吗？周三呢？」"),
    ("🧩", "碎片多", "连堂课中间 10–15 分钟休息，\n手动计算常出现无意义碎片时间"),
    ("🔒", "隐私差", "部分课表 App 需登录上传课表，\n数据存到第三方服务器"),
    ("📵", "依赖网", "校园网不稳定时，在线工具\n无法访问，体验中断"),
]
for i, (icon, title, desc) in enumerate(pain_points):
    col = i % 2
    row = i // 2
    left = Inches(0.8 + col * 6.1)
    top = Inches(2.3 + row * 2.4)
    add_rect(s, left, top, Inches(5.7), Inches(2.0), RGBColor(0x25, 0x25, 0x40))
    add_text(s, left + Inches(0.3), top + Inches(0.2), Inches(1), Inches(0.8),
             icon, size=32, align=PP_ALIGN.LEFT)
    add_text(s, left + Inches(1.2), top + Inches(0.3), Inches(4), Inches(0.5),
             title, size=20, color=ACCENT, bold=True)
    add_text(s, left + Inches(1.2), top + Inches(0.9), Inches(4.3), Inches(1.0),
             desc, size=13, color=LIGHT)
page_number(s, 3, TOTAL)

# ============ Slide 4: 解决方案 ============
s = prs.slides.add_slide(blank)
set_bg(s, BG_DARK)
add_accent_bar(s, Inches(0.6), Inches(0.7))
add_text(s, Inches(0.9), Inches(0.6), Inches(10), Inches(0.7),
         "03  解决方案", size=32, color=WHITE, bold=True)
add_text(s, Inches(0.9), Inches(1.5), Inches(11), Inches(0.5),
         "校园工具箱 — 一个聚合、纯前端、离线可用的工具平台", size=18, color=ACCENT)

# 左侧：架构图
add_rect(s, Inches(0.8), Inches(2.3), Inches(6.2), Inches(4.5), RGBColor(0x25, 0x25, 0x40))
add_text(s, Inches(1.1), Inches(2.5), Inches(5.8), Inches(0.5),
         "技术架构", size=18, color=ACCENT, bold=True)

layers = [
    ("校园专属工具层", "课表空闲计算 · 绩点换算 · 公文自检 · 文献格式转换", ACCENT),
    ("通用工具层", "二维码 · 哈希 · JSON · 编码转换 · 加密 …（来自 it-tools）", ACCENT2),
    ("底座框架", "Vue 3 + TypeScript + Vite + Naive UI", LIGHT),
    ("PWA 离线层", "Service Worker 缓存，断网可用", WHITE),
]
for i, (name, desc, color) in enumerate(layers):
    top = Inches(3.2 + i * 0.85)
    add_rect(s, Inches(1.1), top, Inches(5.6), Inches(0.7), RGBColor(0x30, 0x30, 0x50))
    add_text(s, Inches(1.3), top + Inches(0.1), Inches(2.2), Inches(0.5),
             name, size=14, color=color, bold=True)
    add_text(s, Inches(3.4), top + Inches(0.12), Inches(3.2), Inches(0.5),
             desc, size=12, color=LIGHT)

# 右侧：四大特性
features = [
    ("纯前端计算", "所有计算在浏览器完成，数据不上传"),
    ("全中文界面", "i18n 本地化，符合使用习惯"),
    ("离线可用", "PWA 缓存，校园网断了也能用"),
    ("开源可扩展", "一个工具=一个组件，人人可贡献"),
]
for i, (t, d) in enumerate(features):
    top = Inches(2.3 + i * 1.15)
    add_rect(s, Inches(7.3), top, Inches(5.3), Inches(0.95), RGBColor(0x25, 0x25, 0x40))
    add_accent_bar(s, Inches(7.3), top + Inches(0.25), Inches(0.08), Inches(0.45))
    add_text(s, Inches(7.55), top + Inches(0.12), Inches(4.8), Inches(0.4),
             t, size=16, color=ACCENT, bold=True)
    add_text(s, Inches(7.55), top + Inches(0.5), Inches(4.9), Inches(0.4),
             d, size=12, color=LIGHT)
page_number(s, 4, TOTAL)

# ============ Slide 5: 核心功能 ============
s = prs.slides.add_slide(blank)
set_bg(s, BG_DARK)
add_accent_bar(s, Inches(0.6), Inches(0.7))
add_text(s, Inches(0.9), Inches(0.6), Inches(10), Inches(0.7),
         "04  核心功能", size=32, color=WHITE, bold=True)
add_text(s, Inches(0.9), Inches(1.5), Inches(11), Inches(0.5),
         "本次新增：课表空闲时段计算", size=18, color=ACCENT)

funcs = [
    ("① 读入课表", "粘贴 CSV：课程名,星期几,开始,结束\n自动解析，支持注释行"),
    ("② 单人空闲", "设定每日可用范围（默认 08:00–22:00）\n输出每天的空闲时段"),
    ("③ 碎片合并", "连堂课小间隔（默认 ≤15 分钟）自动合并\n消除 10 分钟无意义碎片"),
    ("④ 共同空闲", "输入多人课表，求交集\n按时长从长到短排序，快速找到约见时间"),
]
for i, (t, d) in enumerate(funcs):
    col = i % 2
    row = i // 2
    left = Inches(0.8 + col * 6.1)
    top = Inches(2.3 + row * 2.4)
    add_rect(s, left, top, Inches(5.7), Inches(2.0), RGBColor(0x25, 0x25, 0x40))
    add_text(s, left + Inches(0.4), top + Inches(0.3), Inches(5), Inches(0.5),
             t, size=20, color=ACCENT, bold=True)
    add_text(s, left + Inches(0.4), top + Inches(0.9), Inches(5), Inches(0.9),
             d, size=13, color=LIGHT)
page_number(s, 5, TOTAL)

# ============ Slide 6: 竞品对比 ============
s = prs.slides.add_slide(blank)
set_bg(s, BG_DARK)
add_accent_bar(s, Inches(0.6), Inches(0.7))
add_text(s, Inches(0.9), Inches(0.6), Inches(10), Inches(0.7),
         "05  竞争优势", size=32, color=WHITE, bold=True)
add_text(s, Inches(0.9), Inches(1.5), Inches(11), Inches(0.5),
         "与通用课表 App / 在线工具对比", size=18, color=ACCENT)

# 表格
headers = ["对比项", "通用课表 App", "在线课表工具", "校园工具箱（本项目）"]
rows = [
    ["共同空闲计算", "✗ 仅展示个人", "部分支持", "✓ 多人共同空闲+排序"],
    ["碎片合并", "✗", "✗", "✓ 可配置阈值自动合并"],
    ["数据隐私", "需登录上传", "可能上传", "✓ 纯前端，不上传"],
    ["离线可用", "部分支持", "✗ 依赖网络", "✓ PWA 离线缓存"],
    ["开源可审计", "✗ 闭源", "部分", "✓ GPL-3.0 完全开源"],
    ["校园定制", "✗", "✗", "✓ 面向大学生场景"],
]
# 表头
col_w = [Inches(2.6), Inches(2.8), Inches(2.8), Inches(4.2)]
x0 = Inches(0.8)
y0 = Inches(2.3)
for j, h in enumerate(headers):
    left = x0 + sum(col_w[:j], Emu(0))
    add_rect(s, left, y0, col_w[j], Inches(0.55), ACCENT2)
    add_text(s, left, y0 + Inches(0.1), col_w[j], Inches(0.4),
             h, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
for i, row in enumerate(rows):
    top = y0 + Inches(0.55) + Inches(i * 0.62)
    for j, cell in enumerate(row):
        left = x0 + sum(col_w[:j], Emu(0))
        bg = RGBColor(0x25, 0x25, 0x40) if i % 2 == 0 else RGBColor(0x2D, 0x2D, 0x48)
        if j == 3:
            bg = RGBColor(0x1E, 0x3A, 0x40)
        add_rect(s, left, top, col_w[j], Inches(0.62), bg)
        color = ACCENT if j == 3 else LIGHT
        bold = j == 3
        add_text(s, left + Inches(0.1), top + Inches(0.14), col_w[j] - Inches(0.2), Inches(0.4),
                 cell, size=12, color=color, bold=bold, align=PP_ALIGN.CENTER)
page_number(s, 6, TOTAL)

# ============ Slide 7: 技术亮点 ============
s = prs.slides.add_slide(blank)
set_bg(s, BG_DARK)
add_accent_bar(s, Inches(0.6), Inches(0.7))
add_text(s, Inches(0.9), Inches(0.6), Inches(10), Inches(0.7),
         "06  技术亮点", size=32, color=WHITE, bold=True)

techs = [
    ("时间建模", "统一用「从 00:00 起的分钟数」表示\n避免字符串比较出错，区间运算高效"),
    ("区间合并", "排序后贪心合并，gap_tol 吸收课间小间隔\nO(n log n)，连堂课不再产生碎片"),
    ("区间交集", "双指针求多人空闲区间交集\n得到共同空闲时段"),
    ("PWA 离线", "基于 vite-plugin-pwa 的 Service Worker\n静态资源预缓存，断网可访问"),
    ("纯前端", "无后端服务，所有计算在浏览器完成\n零服务器成本，数据隐私有保障"),
    ("组件化", "遵循 it-tools「一工具一组件」规范\n新增工具门槛低，便于学生共建"),
]
for i, (t, d) in enumerate(techs):
    col = i % 3
    row = i // 3
    left = Inches(0.8 + col * 4.1)
    top = Inches(2.0 + row * 2.5)
    add_rect(s, left, top, Inches(3.7), Inches(2.1), RGBColor(0x25, 0x25, 0x40))
    add_accent_bar(s, left + Inches(0.3), top + Inches(0.3), Inches(0.08), Inches(0.5))
    add_text(s, left + Inches(0.5), top + Inches(0.25), Inches(3), Inches(0.5),
             t, size=18, color=ACCENT, bold=True)
    add_text(s, left + Inches(0.5), top + Inches(0.85), Inches(3), Inches(1.1),
             d, size=12, color=LIGHT)
page_number(s, 7, TOTAL)

# ============ Slide 8: 推广方案 ============
s = prs.slides.add_slide(blank)
set_bg(s, BG_DARK)
add_accent_bar(s, Inches(0.6), Inches(0.7))
add_text(s, Inches(0.9), Inches(0.6), Inches(10), Inches(0.7),
         "07  推广方案", size=32, color=WHITE, bold=True)
add_text(s, Inches(0.9), Inches(1.5), Inches(11), Inches(0.5),
         "怎么让同学用起来", size=18, color=ACCENT)

steps = [
    ("1", "班级试点", "先在本班/本院推广\n收集反馈快速迭代"),
    ("2", "场景绑定", "约自习 · 组队作业 · 社团排会\n在高频场景下推荐"),
    ("3", "鼓励共建", "架构简单，学生可贡献工具\n形成正向生态循环"),
    ("4", "校内部署", "部署到学校服务器/静态托管\n提供稳定访问地址"),
]
for i, (num, t, d) in enumerate(steps):
    left = Inches(0.8 + i * 3.1)
    top = Inches(2.5)
    # 圆圈数字
    circle = s.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(1.0), top, Inches(1.0), Inches(1.0))
    circle.fill.solid(); circle.fill.fore_color.rgb = ACCENT; circle.line.fill.background()
    add_text(s, left + Inches(1.0), top + Inches(0.2), Inches(1.0), Inches(0.6),
             num, size=32, color=BG_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, left, top + Inches(1.2), Inches(3.0), Inches(0.5),
             t, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, left, top + Inches(1.8), Inches(3.0), Inches(1.2),
             d, size=12, color=LIGHT, align=PP_ALIGN.CENTER)

# 后续方向
add_text(s, Inches(0.9), Inches(5.3), Inches(11), Inches(0.5),
         "后续可扩展工具方向", size=18, color=ACCENT, bold=True)
add_rect(s, Inches(0.8), Inches(5.9), Inches(11.7), Inches(1.0), RGBColor(0x25, 0x25, 0x40))
add_text(s, Inches(1.0), Inches(6.1), Inches(11.3), Inches(0.7),
         "绩点/加权平均分换算  ·  课表冲突检测  ·  公文格式自检  ·  论文参考文献格式转换  ·  考试倒计时",
         size=14, color=LIGHT, align=PP_ALIGN.CENTER)
page_number(s, 8, TOTAL)

# ============ Slide 9: GPL 许可证说明 ============
s = prs.slides.add_slide(blank)
set_bg(s, BG_DARK)
add_accent_bar(s, Inches(0.6), Inches(0.7))
add_text(s, Inches(0.9), Inches(0.6), Inches(10), Inches(0.7),
         "08  开源许可证", size=32, color=WHITE, bold=True)
add_text(s, Inches(0.9), Inches(1.5), Inches(11), Inches(0.5),
         "对 GPL-3.0 「传染性」的理解与遵守", size=18, color=ACCENT)

# 左侧说明
add_rect(s, Inches(0.8), Inches(2.3), Inches(7.2), Inches(4.5), RGBColor(0x25, 0x25, 0x40))
add_text(s, Inches(1.1), Inches(2.5), Inches(6.8), Inches(0.5),
         "什么是 GPL-3.0 的传染性？", size=18, color=ACCENT, bold=True)
add_text(s, Inches(1.1), Inches(3.1), Inches(6.8), Inches(3.5),
         "GPL-3.0 是强 copyleft 许可证。\n\n"
         "核心规则：只要你的项目使用了（链接/修改/衍生）GPL-3.0 代码，\n"
         "并且对外发布（分发），那么整个衍生作品也必须以 GPL-3.0 开源。\n\n"
         "对本项目的影响：\n"
         "• it-tools 采用 GPL-3.0，我们 fork 并修改后发布给同学使用，\n"
         "  因此所有改动（含新增工具）必须以 GPL-3.0 开源。\n"
         "• 不能将改造版本变成闭源商业产品。\n"
         "• 我们的 fork 已保留 LICENSE 文件，符合 GPL-3.0 要求。",
         size=13, color=LIGHT)

# 右侧承诺
add_rect(s, Inches(8.3), Inches(2.3), Inches(4.2), Inches(4.5), ACCENT2)
add_text(s, Inches(8.5), Inches(2.7), Inches(3.8), Inches(0.5),
         "我们的承诺", size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
commitments = [
    "✓ 完全开源，GPL-3.0",
    "✓ 保留原 LICENSE",
    "✓ 欢迎自由使用/修改",
    "✓ 不用于闭源商业",
    "✓ 分发时保持开源",
]
for i, c in enumerate(commitments):
    add_text(s, Inches(8.6), Inches(3.5 + i * 0.6), Inches(3.8), Inches(0.5),
             c, size=15, color=WHITE, align=PP_ALIGN.CENTER)
page_number(s, 9, TOTAL)

# ============ Slide 10: 项目成果 ============
s = prs.slides.add_slide(blank)
set_bg(s, BG_DARK)
add_accent_bar(s, Inches(0.6), Inches(0.7))
add_text(s, Inches(0.9), Inches(0.6), Inches(10), Inches(0.7),
         "09  项目成果", size=32, color=WHITE, bold=True)

results = [
    ("✓", "Fork it-tools 并新增工具"),
    ("✓", "课表空闲时段计算（原项目没有）"),
    ("✓", "中文界面与本地化"),
    ("✓", "PWA 离线可用"),
    ("✓", "本地构建验证通过"),
    ("✓", "项目策划书"),
    ("✓", "双创大赛 PPT"),
    ("✓", "GPL-3.0 合规"),
]
for i, (mark, text) in enumerate(results):
    col = i % 2
    row = i // 2
    left = Inches(1.0 + col * 5.8)
    top = Inches(2.0 + row * 0.85)
    add_text(s, left, top, Inches(0.6), Inches(0.5),
             mark, size=24, color=ACCENT, bold=True)
    add_text(s, left + Inches(0.7), top + Inches(0.05), Inches(5), Inches(0.5),
             text, size=16, color=WHITE)
page_number(s, 10, TOTAL)

# ============ Slide 11: 结尾 ============
s = prs.slides.add_slide(blank)
set_bg(s, BG_DARK)
c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4), Inches(1), Inches(5), Inches(5))
c.fill.solid(); c.fill.fore_color.rgb = ACCENT2; c.line.fill.background()
c.fill.fore_color.brightness = 0.3
add_text(s, Inches(0), Inches(2.8), Inches(13.333), Inches(1.2),
         "感谢聆听", size=64, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(s, Inches(0), Inches(4.2), Inches(13.333), Inches(0.6),
         "校园工具箱 · 让校园生活更高效", size=22, color=ACCENT, align=PP_ALIGN.CENTER)
add_text(s, Inches(0), Inches(5.2), Inches(13.333), Inches(0.5),
         "开源地址：https://github.com/eleven1130/-", size=14, color=GRAY, align=PP_ALIGN.CENTER)
page_number(s, 11, TOTAL)

out = "/workspace/it-tools/校园工具箱_双创大赛PPT.pptx"
prs.save(out)
print(f"PPT 已生成: {out}")
print(f"共 {len(prs.slides)} 页")
